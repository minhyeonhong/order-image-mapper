from openpyxl import load_workbook


class ExcelHandler:
    def __init__(self, filepath):
        self.filepath = filepath
        self.wb = load_workbook(filepath)
        self.ws = self.wb.active

    def read_products(
        self,
        start_row,
        jan_col,
        name_col,
    ):

        products = []

        for row in range(start_row, self.ws.max_row + 1):

            jan = self.ws[f'{jan_col}{row}'].value
            name = self.ws[f'{name_col}{row}'].value

            if not jan and not name:
                continue

            products.append({
                'row': row,
                'jan': str(jan) if jan else '',
                'name': str(name) if name else '',
            })

        return products

    def save(self, path):
        self.wb.save(path)

    def create_failed_sheet(self, failed_items):

        if '매칭실패' in self.wb.sheetnames:
            del self.wb['매칭실패']

        ws = self.wb.create_sheet('매칭실패')

        ws['A1'] = 'JAN'
        ws['B1'] = '상품명'

        for idx, item in enumerate(failed_items, start=2):
            ws[f'A{idx}'] = item['jan']
            ws[f'B{idx}'] = item['name']